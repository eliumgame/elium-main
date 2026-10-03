"""
Regenerates the Office-made part of the « Modifier le texte » corpus
(tests/fixtures/textedit/): Word and PowerPoint documents built with
python-docx / python-pptx, exported to PDF by LibreOffice (embedded subset
TrueType fonts, as real documents have), and a page whose text lives in a
Form XObject (pikepdf). The pdf-lib / Chromium part is made by
tests/support/textedit-fixtures.mjs.

Needs: soffice, the Carlito and Liberation fonts installed, and
python-docx, python-pptx, pikepdf.

    python tests/support/textedit-fixtures.py
"""

import os
import shutil
import subprocess
import sys
import tempfile

import pikepdf
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PptColor
from pptx.util import Pt as PptPt, Cm as PptCm

OUT = os.path.join(os.path.dirname(__file__), "..", "fixtures", "textedit")

LOREM_FR = (
    "Le présent document décrit les conditions dans lesquelles les deux parties entendent "
    "collaborer pendant la durée du projet. Chaque partie conserve la propriété de ses outils "
    "et de ses méthodes, et s'engage à informer l'autre de toute difficulté susceptible de "
    "retarder la livraison des travaux convenus entre elles."
)


def no_space(p, before=0, after=0, line=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line:
        pf.line_spacing = line
    return p


def font(run, name="Carlito", size=11, bold=None, color=None, italic=None):
    run.font.name = name
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    # East-Asian / complex-script slots too, or LibreOffice may pick another face.
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(k), name)
    return run


def letter(path):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Carlito"
    st.font.size = Pt(11)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = Cm(2)
    sec.bottom_margin = Cm(2)
    font(sec.header.paragraphs[0].add_run("Elium — courrier de test"), size=9, color="555555")
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(fp.add_run("Pied de page — document de démonstration"), size=9)

    for line in ("Société Exemple SARL", "12 rue de la Paix", "75002 Paris"):
        font(no_space(doc.add_paragraph()).add_run(line))
    p = no_space(doc.add_paragraph(), before=6, after=6)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    font(p.add_run("Paris, le 3 octobre 2026"))

    p = no_space(doc.add_paragraph(), before=6, after=10)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(p.add_run("Objet : proposition de partenariat"), size=15, bold=True)

    p = no_space(doc.add_paragraph(), after=8)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(1)
    font(p.add_run(LOREM_FR), name="Liberation Serif", size=11)

    p = no_space(doc.add_paragraph(), after=8)
    font(p.add_run("Nous attirons votre attention sur un point "))
    font(p.add_run("important"), bold=True)
    font(p.add_run(" : les délais indiqués en "))
    font(p.add_run("rouge"), color="C00000")
    font(p.add_run(" sont impératifs, ceux indiqués en noir restent indicatifs et pourront être revus d'un commun accord."))

    items = [
        "Une première étape de cadrage du besoin, menée avec les équipes métier pendant deux semaines.",
        "Une phase de réalisation.",
        "Une recette finale suivie d'une mise en production accompagnée.",
    ]
    for it in items:
        p = no_space(doc.add_paragraph(style="List Bullet"))
        font(p.add_run(it))
    for it in ("Signature du contrat.", "Versement de l'acompte prévu au contrat, dans les trente jours suivant la signature."):
        p = no_space(doc.add_paragraph(style="List Number"))
        font(p.add_run(it))

    no_space(doc.add_paragraph(), after=4)
    table = doc.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    for r, row in enumerate([("Désignation", "Quantité", "Prix"), ("Chaise", "4", "45,00"), ("Bureau", "1", "320,00")]):
        for c, txt in enumerate(row):
            cell = table.cell(r, c)
            font(no_space(cell.paragraphs[0]).add_run(txt), bold=(r == 0))

    p = no_space(doc.add_paragraph(), before=10)
    font(p.add_run("Veuillez agréer nos salutations distinguées."))

    # Page 2: a section in two columns.
    new = doc.add_section(WD_SECTION.NEW_PAGE)
    cols = new._sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        new._sectPr.append(cols)
    cols.set(qn("w:num"), "2")
    cols.set(qn("w:space"), "567")
    for i in range(4):
        p = no_space(doc.add_paragraph(), after=6)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if i % 2 else WD_ALIGN_PARAGRAPH.LEFT
        font(p.add_run(f"Colonne, paragraphe {i + 1}. " + LOREM_FR[: 160 + 20 * i]))
        if i == 1:
            br = OxmlElement("w:br")
            br.set(qn("w:type"), "column")
            p.runs[-1]._element.append(br)
    doc.save(path)


def slides(path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = PptCm(25.4), PptCm(14.29)
    s = prs.slides.add_slide(prs.slide_layouts[6])

    def box(x, y, w, h, paras, rot=0):
        tb = s.shapes.add_textbox(PptCm(x), PptCm(y), PptCm(w), PptCm(h))
        tb.rotation = rot
        tf = tb.text_frame
        tf.word_wrap = True
        for i, (text, size, bold, color) in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            r = p.add_run()
            r.text = text
            r.font.name = "Liberation Sans"
            r.font.size = PptPt(size)
            r.font.bold = bold
            if color:
                r.font.color.rgb = PptColor.from_string(color)
        return tb

    box(1.5, 0.8, 22, 2, [("Bilan trimestriel du projet", 30, True, "1F3864")])
    box(
        1.5,
        3.2,
        11,
        8,
        [
            ("Les objectifs du trimestre ont été atteints dans les délais prévus par le plan de charge.", 16, False, None),
            ("Le budget reste maîtrisé.", 16, False, "C00000"),
        ],
    )
    box(14, 4, 9, 3, [("Texte tourné de quinze degrés", 18, True, "2E7D32")], rot=-15)
    prs.save(path)


def to_pdf(src, outdir):
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "pdf", "--outdir", outdir, src],
        check=True,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    return os.path.join(outdir, os.path.splitext(os.path.basename(src))[0] + ".pdf")


def form_xobject(src_pdf, path):
    """Page 1 draws a form holding a LibreOffice page's whole text (offset); page 2 draws the same form."""
    src = pikepdf.open(src_pdf)
    out = pikepdf.new()
    out.pages.append(src.pages[0])
    page = out.pages[0]
    form = page.as_form_xobject()
    form_ref = out.make_indirect(form)
    mb = page.mediabox
    for _ in range(2):
        out.add_blank_page(page_size=(float(mb[2]), float(mb[3])))
    for i, (dx, dy) in ((1, (15, -25)), (2, (0, 0))):
        pg = out.pages[i]
        pg.Resources = pikepdf.Dictionary(XObject=pikepdf.Dictionary(Fm0=form_ref))
        pg.Contents = out.make_stream(f"q 1 0 0 1 {dx} {dy} cm /Fm0 Do Q\n".encode())
    del out.pages[0]
    out.save(path, compress_streams=True, object_stream_mode=pikepdf.ObjectStreamMode.generate)


def main():
    os.makedirs(OUT, exist_ok=True)
    tmp = tempfile.mkdtemp()
    try:
        letter(os.path.join(tmp, "letter.docx"))
        slides(os.path.join(tmp, "slides.pptx"))
        short = Document()
        short.styles["Normal"].font.name = "Liberation Serif"
        for t in ("Texte dans un formulaire XObject", LOREM_FR, "Fin du formulaire."):
            font(short.add_paragraph().add_run(t), name="Liberation Serif", size=12)
        short.save(os.path.join(tmp, "formsrc.docx"))
        for name in ("letter", "slides"):
            pdf = to_pdf(os.path.join(tmp, f"{name}.{'pptx' if name == 'slides' else 'docx'}"), tmp)
            shutil.copy(pdf, os.path.join(OUT, f"{name}.pdf"))
        form_xobject(to_pdf(os.path.join(tmp, "formsrc.docx"), tmp), os.path.join(OUT, "form-xobject.pdf"))
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    for f in sorted(os.listdir(OUT)):
        print(f, os.path.getsize(os.path.join(OUT, f)))


if __name__ == "__main__":
    sys.exit(main())
